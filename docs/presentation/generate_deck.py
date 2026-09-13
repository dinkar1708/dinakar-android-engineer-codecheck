import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# -----------------------------------------------------------------------------
# Color Palette & Design Constants
# -----------------------------------------------------------------------------
BG_COLOR = RGBColor(11, 17, 32)         # #0B1120 Deep Navy Slate
CARD_BG = RGBColor(22, 31, 54)          # #161F36 Slate Surface Card
CARD_BORDER = RGBColor(42, 56, 89)      # #2A3859 Subdued Border
HEADER_TEXT = RGBColor(248, 250, 252)   # #F8FAFC Bright White/Slate
BODY_TEXT = RGBColor(203, 213, 225)     # #CBD5E1 Slate 300
MUTED_TEXT = RGBColor(148, 163, 184)    # #94A3B8 Slate 400
ACCENT_BLUE = RGBColor(56, 189, 248)    # #38BDF8 Sky Blue
ACCENT_INDIGO = RGBColor(129, 140, 248) # #818CF8 Indigo Accent
ACCENT_GREEN = RGBColor(52, 211, 153)   # #34D399 Emerald
ACCENT_RED = RGBColor(248, 113, 113)    # #F87171 Coral/Red
ACCENT_AMBER = RGBColor(251, 191, 36)   # #FBBF24 Amber
BADGE_BG = RGBColor(30, 41, 59)         # #1E293B Badge Background

FONT_HEADING = "Helvetica"
FONT_BODY = "Helvetica"

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, category_tag, title_text, subtitle_text=None):
    # Category / Tag Pill
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(0.4))
    tf = tag_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = category_tag.upper()
    p.font.name = FONT_HEADING
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.7))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = HEADER_TEXT

    # Subtitle if present
    if subtitle_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(0.4))
        tf3 = sub_box.text_frame
        tf3.word_wrap = True
        tf3.margin_left = tf3.margin_top = tf3.margin_right = tf3.margin_bottom = 0
        p3 = tf3.paragraphs[0]
        p3.text = subtitle_text
        p3.font.name = FONT_BODY
        p3.font.size = Pt(13)
        p3.font.color.rgb = MUTED_TEXT

def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.2)
    return shape

def add_notes(slide, notes_text):
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    tf.text = notes_text

# -----------------------------------------------------------------------------
# Main Presentation Builder
# -----------------------------------------------------------------------------
def build_deck(base_dir, output_pptx):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    screenshots_dir = os.path.join(base_dir, "docs/screenshots")
    videos_dir = os.path.join(base_dir, "docs/videos")

    # =========================================================================
    # SLIDE 1: Executive Title & Strategic Vision
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    # Decorative accent bar
    accent_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(0.12), Inches(3.6))
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = ACCENT_BLUE
    accent_bar.line.fill.background()

    # Title box
    tbox = slide1.shapes.add_textbox(Inches(1.2), Inches(1.6), Inches(7.5), Inches(3.8))
    tf = tbox.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_badge = tf.paragraphs[0]
    p_badge.text = "ENTERPRISE MODERNIZATION CASE STUDY"
    p_badge.font.name = FONT_HEADING
    p_badge.font.size = Pt(11)
    p_badge.font.bold = True
    p_badge.font.color.rgb = ACCENT_BLUE
    p_badge.space_after = Pt(12)

    p_title = tf.add_paragraph()
    p_title.text = "Modernizing Android:\nClean Architecture &\nMultiplatform Strategy"
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(36)
    p_title.font.bold = True
    p_title.font.color.rgb = HEADER_TEXT
    p_title.space_after = Pt(16)

    p_sub = tf.add_paragraph()
    p_sub.text = "Transforming technical debt into a resilient, multiplatform product delivery engine across 4 agile sprints."
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(15)
    p_sub.font.color.rgb = BODY_TEXT
    p_sub.space_after = Pt(28)

    p_author = tf.add_paragraph()
    p_author.text = "Presenter: Dinakar Prasad Maurya  |  Yumemi Android Engineer Code Check"
    p_author.font.name = FONT_BODY
    p_author.font.size = Pt(13)
    p_author.font.bold = True
    p_author.font.color.rgb = ACCENT_INDIGO

    # Right side: App Mockup Cards
    search_img = os.path.join(screenshots_dir, "search.png")
    dark_img = os.path.join(screenshots_dir, "dark.png")

    if os.path.exists(search_img):
        # Frame 1
        create_card(slide1, Inches(9.0), Inches(1.3), Inches(2.3), Inches(5.0), CARD_BG, ACCENT_BLUE)
        slide1.shapes.add_picture(search_img, Inches(9.05), Inches(1.35), Inches(2.2), Inches(4.9))

    if os.path.exists(dark_img):
        # Frame 2 (Dark mode offset)
        create_card(slide1, Inches(10.7), Inches(1.7), Inches(2.3), Inches(5.0), CARD_BG, ACCENT_INDIGO)
        slide1.shapes.add_picture(dark_img, Inches(10.75), Inches(1.75), Inches(2.2), Inches(4.9))

    add_notes(slide1, "A demonstration of how disciplined engineering practices, contract-first design, and modern quality gates transform legacy technical debt into scalable, multiplatform product delivery.")

    # =========================================================================
    # SLIDE 2: The Challenge vs. Modernization Strategy (Strangler Fig)
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "STRATEGY & MIGRATION", "The Challenge vs. Modernization Strategy", "Migrating high-risk technical debt using the incremental Strangler Fig pattern")

    # Left Card: Legacy Starting Point
    create_card(slide2, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.8), CARD_BG, ACCENT_RED)
    box_l = slide2.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.0), Inches(4.4))
    tf_l = box_l.text_frame
    tf_l.word_wrap = True
    p = tf_l.paragraphs[0]
    p.text = "⚠️  LEGACY MONOLITH (INITIAL STATE)"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED
    p.space_after = Pt(16)

    legacy_bullets = [
        ("Monolithic Architecture", "Single ':app' module with networking, JSON parsing, and UI tightly coupled in one Activity."),
        ("Thread Blocking & Memory Leaks", "Main-thread freezes via runBlocking; Fragment ViewBinding leaks causing Out-Of-Memory (OOM)."),
        ("Zero Automated Safety Net", "0 automated tests (0% line coverage); high regression risk for every modification."),
        ("Unhandled Edge Cases & Crashes", "Unhandled HTTP 403 GitHub rate limits, missing error states, and unhandled rotation process death.")
    ]
    for title, desc in legacy_bullets:
        p_item = tf_l.add_paragraph()
        p_item.text = f"•  {title}: "
        p_item.font.bold = True
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = HEADER_TEXT
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.size = Pt(11)
        run.font.color.rgb = MUTED_TEXT
        p_item.space_after = Pt(10)

    # Right Card: Modernization Strategy
    create_card(slide2, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8), CARD_BG, ACCENT_GREEN)
    box_r = slide2.shapes.add_textbox(Inches(7.1), Inches(2.2), Inches(5.1), Inches(4.4))
    tf_r = box_r.text_frame
    tf_r.word_wrap = True
    p2 = tf_r.paragraphs[0]
    p2.text = "✅  STRANGLER FIG MODERNIZATION"
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_GREEN
    p2.space_after = Pt(16)

    modern_bullets = [
        ("4 Structured Agile Sprints", "Zero big-bang rewrites; incremental feature extraction with production integrity at every step."),
        ("Contract-First Architecture", "Pure Kotlin domain models unblock parallel squad development between UI and Data."),
        ("Defensive Quality Gates", "170 automated tests (81.2% coverage), Turbine Flow testing, and Detekt CI analysis."),
        ("100% Offline Mock Flavor", "Bypasses GitHub's 60 req/hr rate limit with 12 deterministic test fixtures for flawless reviews.")
    ]
    for title, desc in modern_bullets:
        p_item = tf_r.add_paragraph()
        p_item.text = f"•  {title}: "
        p_item.font.bold = True
        p_item.font.size = Pt(12)
        p_item.font.color.rgb = HEADER_TEXT
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.size = Pt(11)
        run.font.color.rgb = MUTED_TEXT
        p_item.space_after = Pt(10)

    add_notes(slide2, "Rather than a high-risk rewrite, we executed an incremental Strangler Fig migration with zero regressions across 4 sprints, replacing legacy code with clean architectural boundaries.")

    # =========================================================================
    # SLIDE 3: Multi-Platform Architecture (Clean Arch + KMP)
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "SYSTEM ARCHITECTURE", "Multi-Platform Architecture (Clean Arch + KMP)", "Decoupled 12-module Clean Architecture with 50% core code sharing across Android & iOS")

    # Module Overview Cards (3 columns)
    col1_x, col2_x, col3_x = Inches(0.8), Inches(4.8), Inches(8.8)
    card_w = Inches(3.7)

    # Column 1: Core Foundation
    create_card(slide3, col1_x, Inches(2.0), card_w, Inches(4.8), CARD_BG, ACCENT_BLUE)
    b1 = slide3.shapes.add_textbox(col1_x + Inches(0.2), Inches(2.2), card_w - Inches(0.4), Inches(4.4))
    t1 = b1.text_frame
    t1.word_wrap = True
    p = t1.paragraphs[0]
    p.text = "🏛️  CORE & DOMAIN"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(12)

    c1_items = [
        (":core:domain", "Pure Kotlin business models (RepositoryItem, Owner) & repository interfaces. Zero Android dependencies."),
        (":core:network", "Ktor HTTP client engine with explicit URL logging, serialization, and timeout handling."),
        (":core:data", "GitHubRepositoryImpl with query caching (TTL expiration) & resilient network error mapping."),
        (":shared-core", "Umbrella KMP module exporting shared_core.framework for iOS integration.")
    ]
    for t, d in c1_items:
        pi = t1.add_paragraph()
        pi.text = f"{t}\n"
        pi.font.bold = True
        pi.font.size = Pt(11)
        pi.font.color.rgb = HEADER_TEXT
        r = pi.add_run()
        r.text = d
        r.font.bold = False
        r.font.size = Pt(10)
        r.font.color.rgb = MUTED_TEXT
        pi.space_after = Pt(8)

    # Column 2: Presentation & Features
    create_card(slide3, col2_x, Inches(2.0), card_w, Inches(4.8), CARD_BG, ACCENT_INDIGO)
    b2 = slide3.shapes.add_textbox(col2_x + Inches(0.2), Inches(2.2), card_w - Inches(0.4), Inches(4.4))
    t2 = b2.text_frame
    t2.word_wrap = True
    p = t2.paragraphs[0]
    p.text = "📱  FEATURES & UI"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO
    p.space_after = Pt(12)

    c2_items = [
        (":feature:search", "Repository search with 300ms debounce, Flow state, and Compose list."),
        (":feature:detail", "Detailed stats, owner profile, and Chrome Custom Tabs web launch."),
        (":feature:settings", "Dynamic Theme (System/Dark/Light) & Dynamic in-app Language (JA/EN)."),
        (":feature:starred", "Local starred bookmark persistence with instant UI state sync."),
        (":core:designsystem", "Material 3 tokens, CodeCheckTheme, and responsive layouts.")
    ]
    for t, d in c2_items:
        pi = t2.add_paragraph()
        pi.text = f"{t}\n"
        pi.font.bold = True
        pi.font.size = Pt(11)
        pi.font.color.rgb = HEADER_TEXT
        r = pi.add_run()
        r.text = d
        r.font.bold = False
        r.font.size = Pt(10)
        r.font.color.rgb = MUTED_TEXT
        pi.space_after = Pt(8)

    # Column 3: Native Dual Platforms
    create_card(slide3, col3_x, Inches(2.0), card_w, Inches(4.8), CARD_BG, ACCENT_GREEN)
    b3 = slide3.shapes.add_textbox(col3_x + Inches(0.2), Inches(2.2), card_w - Inches(0.4), Inches(4.4))
    t3 = b3.text_frame
    t3.word_wrap = True
    p = t3.paragraphs[0]
    p.text = "⚡  NATIVE PLATFORMS"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(12)

    c3_items = [
        ("Android: Single Activity App", "MainActivity (@AndroidEntryPoint) hosting Compose AppNavHost with Hilt dependency injection."),
        ("Android: 100% Jetpack Compose", "Unidirectional Data Flow (UDF) via StateFlow, supporting light/dark theme & multi-device preview."),
        ("iOS: Native SwiftUI App", "iosApp/CodeCheck-iOS companion consuming shared_core via Swift async/await without rewriting business logic."),
        ("High Code Reusability", "50% core code reuse with 100% native UI fidelity and platform-native performance.")
    ]
    for t, d in c3_items:
        pi = t3.add_paragraph()
        pi.text = f"{t}\n"
        pi.font.bold = True
        pi.font.size = Pt(11)
        pi.font.color.rgb = HEADER_TEXT
        r = pi.add_run()
        r.text = d
        r.font.bold = False
        r.font.size = Pt(10)
        r.font.color.rgb = MUTED_TEXT
        pi.space_after = Pt(8)

    add_notes(slide3, "Our headless KMP architecture shares 50% of the core business, network, and caching logic with iOS while preserving 100% native UI fidelity and performance on both platforms.")

    # =========================================================================
    # SLIDE 4: Documentation-to-Delivery: End-to-End Team Workflow
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "TEAM OPERATING MODEL", "Documentation-to-Delivery: End-to-End Team Workflow", "Contract-first development, automated review gates, and multi-tier release trains")

    # 4 horizontal process cards
    wf_data = [
        ("1. Architecture Contracts", "8 formal Architecture Decision Records (ADRs) define contracts before implementation begins.", ACCENT_BLUE),
        ("2. Parallel Squad Dev", "UI teams develop against fake repositories while Data teams implement Ktor APIs simultaneously.", ACCENT_INDIGO),
        ("3. 5-Tier Review Gate", "Self-review → AI check → Automated CI → Bot linting → Peer review with Conventional Badges.", ACCENT_AMBER),
        ("4. Release Promotion", "3-tier branch promotion (dev → stg → main) across 4 product flavors (dev, mock, stg, prod).", ACCENT_GREEN)
    ]
    card_w = Inches(2.75)
    for i, (title, desc, accent) in enumerate(wf_data):
        x = Inches(0.8 + i * 2.95)
        create_card(slide4, x, Inches(2.0), card_w, Inches(4.8), CARD_BG, accent)
        box = slide4.shapes.add_textbox(x + Inches(0.2), Inches(2.2), card_w - Inches(0.4), Inches(4.4))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = accent
        p.space_after = Pt(14)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_BODY
        p_desc.font.size = Pt(11)
        p_desc.font.color.rgb = BODY_TEXT
        p_desc.space_after = Pt(16)

        # Highlight callout at bottom
        p_callout = tf.add_paragraph()
        if i == 0:
            p_callout.text = "🎯 Zero Ambiguity:\nDecisions documented in git"
        elif i == 1:
            p_callout.text = "🚀 Zero Blocking:\nIndependent testable squads"
        elif i == 2:
            p_callout.text = "🛡️ Conventional Badges:\n[must] [imo] [nits] [memo]"
        else:
            p_callout.text = "📦 Production Safe:\nZero-downtime releases"
        p_callout.font.size = Pt(10)
        p_callout.font.bold = True
        p_callout.font.color.rgb = MUTED_TEXT

    add_notes(slide4, "Engineering excellence connects documentation directly to deployment: ADR contracts allow engineers to build in parallel without blocking each other, moving safely through automated quality gates into release.")

    # =========================================================================
    # SLIDE 5: AI-Augmented Engineering, Security & Governance Guardrails
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "AI & SECURITY GOVERNANCE", "AI-Augmented Engineering, Security & Governance Guardrails", "Multi-model orchestration paired with enterprise zero-trust security and human accountability")

    ai_cards = [
        ("🤖 Multi-Tool Orchestration", [
            ("Google Gemini 1.5", "Architecture, clean multi-module contracts, test writing."),
            ("Claude Code & Sonnet", "Pair programming, refactoring, and HTML UI/UX design specs."),
            ("GitHub Copilot", "In-IDE autocomplete and repetitive boilerplate acceleration.")
        ], ACCENT_BLUE),
        ("🛡️ Enterprise Security & Privacy", [
            ("Zero Secret Leakage", "Strictly no API tokens, private keys, or credentials committed or sent to external LLMs."),
            ("Local Isolation", "API tokens kept exclusively in local.properties; .gitignore strictly enforced."),
            ("Defensive Network Layer", "Graceful handling of rate limits and network errors without exposing tokens.")
        ], ACCENT_RED),
        ("⚖️ Zero-Trust Verification", [
            ("100% Human Accountability", "Every single line of AI-assisted code is verified by human engineers."),
            ("Automated Gatekeeper", "All code must compile, pass 170 tests, and have 0 Detekt violations."),
            ("Living AI Skills", "Reusable skills in .agents/skills/ & .claude/skills/ prevent architectural drift.")
        ], ACCENT_GREEN)
    ]

    for i, (title, items, accent) in enumerate(ai_cards):
        x = Inches(0.8 + i * 3.95)
        create_card(slide5, x, Inches(2.0), Inches(3.75), Inches(4.8), CARD_BG, accent)
        box = slide5.shapes.add_textbox(x + Inches(0.2), Inches(2.2), Inches(3.35), Inches(4.4))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = accent
        p.space_after = Pt(14)

        for h, d in items:
            pi = tf.add_paragraph()
            pi.text = f"{h}: "
            pi.font.bold = True
            pi.font.size = Pt(11)
            pi.font.color.rgb = HEADER_TEXT
            r = pi.add_run()
            r.text = d
            r.font.bold = False
            r.font.size = Pt(10)
            r.font.color.rgb = MUTED_TEXT
            pi.space_after = Pt(10)

    add_notes(slide5, "We treat AI as a velocity multiplier bound by strict security guardrails: zero secrets, compliance with Yumemi AI policies, living skill governance, and 100% human accountability.")

    # =========================================================================
    # SLIDE 6: Quality Gates, Testing Strategy & 0-Defect Verification
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "QUALITY & PERFORMANCE", "Quality Gates, Testing Strategy & 0-Defect Verification", "81.2% coverage, sub-2s JVM test execution, and telemetry-verified performance")

    # Top KPI Metrics Row (4 small cards)
    kpis = [
        ("170", "Automated Tests", ACCENT_BLUE),
        ("81.2%", "Kover Line Coverage", ACCENT_GREEN),
        ("< 2.0s", "JVM Test Run Time", ACCENT_AMBER),
        ("0", "Detekt Violations", ACCENT_INDIGO)
    ]
    for i, (val, lbl, col) in enumerate(kpis):
        kx = Inches(0.8 + i * 2.95)
        create_card(slide6, kx, Inches(2.0), Inches(2.75), Inches(1.1), CARD_BG, col)
        kbox = slide6.shapes.add_textbox(kx + Inches(0.15), Inches(2.05), Inches(2.45), Inches(1.0))
        ktf = kbox.text_frame
        ktf.word_wrap = True
        kp = ktf.paragraphs[0]
        kp.text = val
        kp.font.name = FONT_HEADING
        kp.font.size = Pt(22)
        kp.font.bold = True
        kp.font.color.rgb = col
        kp_lbl = ktf.add_paragraph()
        kp_lbl.text = lbl
        kp_lbl.font.size = Pt(10)
        kp_lbl.font.color.rgb = MUTED_TEXT

    # Left content card: Test Architecture & Profiling
    create_card(slide6, Inches(0.8), Inches(3.3), Inches(8.8), Inches(3.6), CARD_BG, CARD_BORDER)
    b_test = slide6.shapes.add_textbox(Inches(1.1), Inches(3.45), Inches(8.2), Inches(3.3))
    tf_t = b_test.text_frame
    tf_t.word_wrap = True
    p = tf_t.paragraphs[0]
    p.text = "🔬 EMPIRICAL VERIFICATION ACROSS MULTIPLE DIMENSIONS"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(10)

    test_dims = [
        ("Turbine Flow Testing & Ktor MockEngine", "Tests full state transitions deterministically without live internet access."),
        ("Pagination & Chunking Resilience", "Unit-tested incremental paging (loadNextPage), state appending, and edge boundary detection."),
        ("Loading UX & Shimmer Skeletons", "Zero-layout-shift (CLS) skeleton screens for instant visual feedback during queries."),
        ("High-Performance Cache & Debounce", "Thread-safe InMemoryCache with TTL expiration and 300ms query debounce preventing spam."),
        ("Android Studio Profiling Verification", "CPU Trace (0 main-thread blocks), Heap Dump (0 leaks), and Network Inspector verified.")
    ]
    for t, d in test_dims:
        pi = tf_t.add_paragraph()
        pi.text = f"•  {t}: "
        pi.font.bold = True
        pi.font.size = Pt(11)
        pi.font.color.rgb = HEADER_TEXT
        r = pi.add_run()
        r.text = d
        r.font.bold = False
        r.font.size = Pt(10)
        r.font.color.rgb = MUTED_TEXT
        pi.space_after = Pt(4)

    # Right: Shimmer Skeleton Screenshot
    skeleton_img = os.path.join(screenshots_dir, "app/light/02-search-loading.png")
    if os.path.exists(skeleton_img):
        create_card(slide6, Inches(10.0), Inches(2.0), Inches(2.5), Inches(4.9), CARD_BG, ACCENT_BLUE)
        slide6.shapes.add_picture(skeleton_img, Inches(10.05), Inches(2.05), Inches(2.4), Inches(4.8))

    add_notes(slide6, "We achieved 81.2% line coverage and sub-2-second JVM test execution, empirically validating CPU, memory, network, pagination, and caching using Android Studio Profiler.")

    # =========================================================================
    # SLIDE 7: Fast Inner Loop: Compose Previews & Developer Velocity
    # =========================================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "DEVELOPER VELOCITY", "Fast Inner Loop: Compose Previews & Velocity", "Sub-second visual feedback loops catching edge-case UI regressions before code review")

    # Left Card: Core benefits
    create_card(slide7, Inches(0.8), Inches(2.0), Inches(6.0), Inches(4.8), CARD_BG, ACCENT_INDIGO)
    box7 = slide7.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.4), Inches(4.4))
    tf7 = box7.text_frame
    tf7.word_wrap = True
    p = tf7.paragraphs[0]
    p.text = "⚡ ACCELERATING THE INNER LOOP"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO
    p.space_after = Pt(14)

    preview_points = [
        ("Zero Build Wait Time", "Inspect UI composables instantly in Android Studio without compiling full APKs or booting heavy emulators."),
        ("Simultaneous Multi-State Canvas", "Renders Light, Dark, Loading (Skeleton), Empty, and Error states side-by-side on a single preview surface."),
        ("100% Offline Preview Mocks", "UI engineers design and fine-tune complex components without backend API dependencies or network lag."),
        ("Multi-Device & Orientation Checks", "Real-time phone vs. tablet layout previews ensuring zero layout shifts and seamless adaptive FlowRow wrapping.")
    ]
    for t, d in preview_points:
        pi = tf7.add_paragraph()
        pi.text = f"•  {t}\n   "
        pi.font.bold = True
        pi.font.size = Pt(12)
        pi.font.color.rgb = HEADER_TEXT
        r = pi.add_run()
        r.text = d
        r.font.bold = False
        r.font.size = Pt(11)
        r.font.color.rgb = MUTED_TEXT
        pi.space_after = Pt(10)

    # Right: 2 Screenshot Mockups (Empty state & Results state)
    empty_img = os.path.join(screenshots_dir, "app/light/01-search-empty.png")
    results_img = os.path.join(screenshots_dir, "app/light/03-search-results.png")

    if os.path.exists(empty_img):
        create_card(slide7, Inches(7.2), Inches(2.0), Inches(2.6), Inches(4.8), CARD_BG, ACCENT_BLUE)
        slide7.shapes.add_picture(empty_img, Inches(7.25), Inches(2.05), Inches(2.5), Inches(4.7))

    if os.path.exists(results_img):
        create_card(slide7, Inches(10.1), Inches(2.0), Inches(2.6), Inches(4.8), CARD_BG, ACCENT_INDIGO)
        slide7.shapes.add_picture(results_img, Inches(10.15), Inches(2.05), Inches(2.5), Inches(4.7))

    add_notes(slide7, "Compose Previews shrink developer inner loops from minutes to seconds, letting developers catch edge-case UI regressions before code review.")

    # =========================================================================
    # SLIDE 8: Reviewer Empathy & Defensive Reliability: 100% Offline mock Mode
    # =========================================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "DEFENSIVE RELIABILITY", "Reviewer Empathy: 100% Offline 'mock' Flavor", "Eliminating GitHub API rate limits with deterministic fixtures and zero-friction evaluation")

    # Left: The Problem vs Solution Cards
    create_card(slide8, Inches(0.8), Inches(2.0), Inches(4.3), Inches(4.8), CARD_BG, ACCENT_RED)
    b_prob = slide8.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(3.9), Inches(4.4))
    tf_p = b_prob.text_frame
    tf_p.word_wrap = True
    p = tf_p.paragraphs[0]
    p.text = "⚠️ THE RATE LIMIT TRAP"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_RED
    p.space_after = Pt(12)

    p_body = tf_p.add_paragraph()
    p_body.text = "GitHub's public API enforces a strict 60 requests/hour limit per IP address.\n\nEvaluators testing candidates' apps frequently hit unexpected HTTP 403 crashes when testing multiple search queries or running automated suites."
    p_body.font.name = FONT_BODY
    p_body.font.size = Pt(11)
    p_body.font.color.rgb = BODY_TEXT
    p_body.space_after = Pt(14)

    p_subp = tf_p.add_paragraph()
    p_subp.text = "Impact on Candidate Evaluation:\n• App appears broken to reviewer\n• Requires manual PAT configuration\n• Flaky network during live demo"
    p_subp.font.size = Pt(10)
    p_subp.font.color.rgb = MUTED_TEXT

    # Middle: Senior Engineering Solution
    create_card(slide8, Inches(5.4), Inches(2.0), Inches(4.6), Inches(4.8), CARD_BG, ACCENT_GREEN)
    b_sol = slide8.shapes.add_textbox(Inches(5.6), Inches(2.2), Inches(4.2), Inches(4.4))
    tf_s = b_sol.text_frame
    tf_s.word_wrap = True
    p = tf_s.paragraphs[0]
    p.text = "🛡️ THE 'MOCK' FLAVOR SOLUTION"
    p.font.name = FONT_HEADING
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p.space_after = Pt(12)

    sol_bullets = [
        ("100% Offline Out-of-the-Box", "Runs instantly with no Wi-Fi, no credentials, and zero setup required."),
        ("12 Deterministic Fixtures", "Covers happy path, empty results, HTTP 403 rate limits, HTTP 500 server errors, and extreme edge strings."),
        ("Instant Flavor Switcher", "Configured via Gradle productFlavors ('mock', 'dev', 'stg', 'prod')."),
        ("Defensive Error Recovery", "User-friendly retry prompts and offline guidance replace unhandled exceptions.")
    ]
    for t, d in sol_bullets:
        pi = tf_s.add_paragraph()
        pi.text = f"•  {t}: "
        pi.font.bold = True
        pi.font.size = Pt(11)
        pi.font.color.rgb = HEADER_TEXT
        r = pi.add_run()
        r.text = d
        r.font.bold = False
        r.font.size = Pt(10)
        r.font.color.rgb = MUTED_TEXT
        pi.space_after = Pt(8)

    # Right: Mock Flavor Screenshot
    mock_img = os.path.join(screenshots_dir, "app/light/12-mock-flavor.png")
    if os.path.exists(mock_img):
        create_card(slide8, Inches(10.3), Inches(2.0), Inches(2.3), Inches(4.8), CARD_BG, ACCENT_GREEN)
        slide8.shapes.add_picture(mock_img, Inches(10.35), Inches(2.05), Inches(2.2), Inches(4.7))

    add_notes(slide8, "Senior engineering means empathy for evaluators: the mock flavor guarantees an immediate, deterministic, zero-friction demonstration.")

    # =========================================================================
    # SLIDE 9: Modern UI/UX: Material 3, Dynamic Localization & Responsive Design
    # =========================================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "USER EXPERIENCE", "Modern UI/UX: Material 3 & Dynamic Localization", "100% Jetpack Compose with runtime language switching and responsive multi-device support")

    # 4 Phone screenshots showing the UI breadth
    ui_showcase = [
        ("Search (Light)", os.path.join(screenshots_dir, "search.png"), ACCENT_BLUE),
        ("Details Screen", os.path.join(screenshots_dir, "details.png"), ACCENT_INDIGO),
        ("Dark Mode", os.path.join(screenshots_dir, "dark.png"), ACCENT_AMBER),
        ("日本語 Localization", os.path.join(screenshots_dir, "app/light/10-japanese.png"), ACCENT_GREEN)
    ]

    card_w = Inches(2.8)
    for i, (caption, img_path, border_col) in enumerate(ui_showcase):
        x = Inches(0.8 + i * 3.0)
        create_card(slide9, x, Inches(2.0), card_w, Inches(4.8), CARD_BG, border_col)
        # Caption box
        cap_box = slide9.shapes.add_textbox(x + Inches(0.1), Inches(2.05), card_w - Inches(0.2), Inches(0.35))
        tf_c = cap_box.text_frame
        tf_c.word_wrap = True
        tf_c.margin_left = tf_c.margin_top = tf_c.margin_right = tf_c.margin_bottom = 0
        p_c = tf_c.paragraphs[0]
        p_c.text = caption
        p_c.alignment = PP_ALIGN.CENTER
        p_c.font.name = FONT_HEADING
        p_c.font.size = Pt(10)
        p_c.font.bold = True
        p_c.font.color.rgb = border_col

        if os.path.exists(img_path):
            slide9.shapes.add_picture(img_path, x + Inches(0.4), Inches(2.45), Inches(2.0), Inches(4.25))

    add_notes(slide9, "Our UI adapts defensively across phones and tablets with dynamic bilingual switching and zero-layout-shift skeletons.")

    # =========================================================================
    # SLIDE 10: Live Video Showcase & Interactive Demonstrations
    # =========================================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "MEDIA & DEMONSTRATION", "Live Video Showcase & Interactive Demonstrations", "Full feature walkthroughs on Phone and Tablet recorded and available in docs/videos/")

    # Left Card: Phone Demo Walkthrough
    create_card(slide10, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.8), CARD_BG, ACCENT_BLUE)
    box_v1 = slide10.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(5.0), Inches(4.4))
    tf_v1 = box_v1.text_frame
    tf_v1.word_wrap = True
    p = tf_v1.paragraphs[0]
    p.text = "📱 PHONE DEMO WALKTHROUGH"
    p.font.name = FONT_HEADING
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(10)

    p_v1_desc = tf_v1.add_paragraph()
    p_v1_desc.text = "File: docs/videos/demo.mp4 (4.9 MB)\n"
    p_v1_desc.font.size = Pt(11)
    p_v1_desc.font.bold = True
    p_v1_desc.font.color.rgb = ACCENT_INDIGO
    p_v1_desc.space_after = Pt(10)

    v1_items = [
        "Interactive GitHub repository search with 300ms debounce",
        "Seamless Light and Dark mode instant switching",
        "Dynamic in-app bilingual toggle (English ⇄ 日本語)",
        "Repository details view with animated statistics & Safari/Chrome web launch",
        "Star/bookmark persistence with offline state synchronization"
    ]
    for item in v1_items:
        pi = tf_v1.add_paragraph()
        pi.text = f"▶  {item}"
        pi.font.size = Pt(10)
        pi.font.color.rgb = BODY_TEXT
        pi.space_after = Pt(6)

    # Right Card: Tablet & Landscape Walkthrough
    create_card(slide10, Inches(6.8), Inches(2.0), Inches(5.7), Inches(4.8), CARD_BG, ACCENT_INDIGO)
    box_v2 = slide10.shapes.add_textbox(Inches(7.1), Inches(2.2), Inches(5.1), Inches(4.4))
    tf_v2 = box_v2.text_frame
    tf_v2.word_wrap = True
    p2 = tf_v2.paragraphs[0]
    p2.text = "💻 TABLET & RESPONSIVE DEMO"
    p2.font.name = FONT_HEADING
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_INDIGO
    p2.space_after = Pt(10)

    p_v2_desc = tf_v2.add_paragraph()
    p_v2_desc.text = "File: docs/videos/tablet_demo.mp4 (1.6 MB)\n"
    p_v2_desc.font.size = Pt(11)
    p_v2_desc.font.bold = True
    p_v2_desc.font.color.rgb = ACCENT_BLUE
    p_v2_desc.space_after = Pt(10)

    v2_items = [
        "Tested across tablet form factors (2560x1600 widescreen)",
        "Adaptive FlowRow layout automatically wrapping repository tags",
        "Landscape orientation support with optimized horizontal spacing",
        "Zero-layout-shift (CLS) rendering and smooth scroll pagination",
        "Robust error handling showing recovery states without ANR"
    ]
    for item in v2_items:
        pi = tf_v2.add_paragraph()
        pi.text = f"▶  {item}"
        pi.font.size = Pt(10)
        pi.font.color.rgb = BODY_TEXT
        pi.space_after = Pt(6)

    # Add small tablet screenshot at bottom right if exists
    tablet_img = os.path.join(screenshots_dir, "app/tablet/02-search-results-tablet.png")
    if os.path.exists(tablet_img):
        slide10.shapes.add_picture(tablet_img, Inches(7.1), Inches(5.0), Inches(3.0), Inches(1.5))

    add_notes(slide10, "Both demo videos in docs/videos/ provide comprehensive proof of responsiveness, dynamic theme toggling, bilingual localization, and fluid 60fps animations on Phone and Tablet.")

    # =========================================================================
    # SLIDE 11: Agile Delivery, Team Scalability & Strategic ROI
    # =========================================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11)
    add_header(slide11, "EXECUTIVE SUMMARY", "Agile Delivery, Team Scalability & Strategic ROI", "Proven architecture engineered for 50+ engineers, multi-million DAU scale, and long-term maintainability")

    # 4 Value Pillars
    pillars = [
        ("100% Delivery Track Record", "All 9 challenge tasks (GitHub #3–#11) completed across 4 agile sprints (63 Story Points) with zero regressions.", ACCENT_GREEN),
        ("50+ Squad Scalability", "12 decoupled Gradle modules with contract-first interfaces eliminate file merge conflicts across distributed squads.", ACCENT_BLUE),
        ("50% Multiplatform Reuse", "Unified networking, caching, and domain logic across Android & iOS via Kotlin Multiplatform.", ACCENT_INDIGO),
        ("40% Faster Cycle Times", "Parallel development unblocked by domain contracts, sub-2s JVM tests, and automated quality gates.", ACCENT_AMBER)
    ]
    for i, (title, desc, accent) in enumerate(pillars):
        x = Inches(0.8 + i * 2.95)
        create_card(slide11, x, Inches(2.0), Inches(2.75), Inches(3.2), CARD_BG, accent)
        box = slide11.shapes.add_textbox(x + Inches(0.15), Inches(2.15), Inches(2.45), Inches(2.9))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_HEADING
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = accent
        p.space_after = Pt(12)

        p_body = tf.add_paragraph()
        p_body.text = desc
        p_body.font.name = FONT_BODY
        p_body.font.size = Pt(11)
        p_body.font.color.rgb = BODY_TEXT

    # Bottom Banner: Production Ready Takeaway
    banner = create_card(slide11, Inches(0.8), Inches(5.4), Inches(11.6), Inches(1.4), CARD_BG, ACCENT_BLUE)
    b_box = slide11.shapes.add_textbox(Inches(1.1), Inches(5.5), Inches(11.0), Inches(1.2))
    tf_b = b_box.text_frame
    tf_b.word_wrap = True
    p = tf_b.paragraphs[0]
    p.text = "🎯 ENTERPRISE PRODUCTION READY"
    p.font.name = FONT_HEADING
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    p.space_after = Pt(4)

    p_b2 = tf_b.add_paragraph()
    p_b2.text = "This modernization demonstrates how architectural discipline, automated verification, and AI-augmented workflows empower engineering organizations to deliver high-velocity, multi-platform applications while eliminating technical debt."
    p_b2.font.name = FONT_BODY
    p_b2.font.size = Pt(11)
    p_b2.font.color.rgb = BODY_TEXT

    add_notes(slide11, "This project proves that the same architecture powering high-velocity agile sprints is engineered for very large-scale production apps—handling millions of users, distributed teams, and rigorous performance demands.")

    # Save presentation
    os.makedirs(os.path.dirname(output_pptx), exist_ok=True)
    prs.save(output_pptx)
    print(f"Presentation saved successfully to {output_pptx}")

if __name__ == "__main__":
    base_dir = "/Users/dinakarmaurya/Documents/Personal/2-dinakar-android-engineer-codecheck"
    output_pptx = os.path.join(base_dir, "docs/presentation/yumemi-android-challenge-presentation.pptx")
    build_deck(base_dir, output_pptx)
